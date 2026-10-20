"""
vacoda.scripts.common.trace_hs_journey

VACODA HS6 lineage diagnostic.

Run from the VACODA project root:

    python -m scripts.common.trace_hs_journey 300610 540600

The script writes a Markdown report that explains the journey of each HS6
code from HS-sync through Validation, including why row counts change.
"""

from __future__ import annotations
import argparse
from datetime import datetime
from pathlib import Path
from typing import Iterable
import duckdb
import pandas as pd
from openpyxl import load_workbook

STAGES = {
    "hs_sync": Path(
        "/storage/vacoda/data/interim/hs_sync/target_h5/applied/"
        "comtrade_target_h5_hs_synced.parquet"
    ),

    "quantity": Path(
        "/storage/vacoda/data/interim/quantity/target_h5/"
        "comtrade_target_h5_quantity_synced.parquet"
    ),

    "cif_fob": Path(
        "/storage/vacoda/data/interim/cif_fob/target_h5/"
        "comtrade_target_h5_cif_fob_corrected.parquet"
    ),

    "trade_sync": Path(
        "/storage/vacoda/data/interim/trade_sync/target_h5/"
        "comtrade_target_h5_trade_synced.parquet"
    ),

    "validation_prepared": Path(
        "/storage/vacoda/data/processed/validation/target_h5/prepared/"
        "comtrade_target_h5_prepared.parquet"
    ),

    "validation_quantity": Path(
        "/storage/vacoda/data/processed/validation/target_h5/quantity/"
        "comtrade_target_h5_quantity_selection.parquet"
    ),
}

STAGE_LABELS = {
    "hs_sync": "HS-sync",
    "quantity": "Quantity",
    "cif_fob": "CIF/FOB",
    "trade_sync": "Trade Sync",
    "validation_prepared": "Validation prepared",
    "validation_quantity": "Validation quantity selection",
}

DEFAULT_REPORT_DIR = Path("/storage/vacoda/output/reports/diagnostics/hs_journey")


# Official Comtrade correlation tables (relative to VACODA project root).
DEFAULT_CORRELATION_DIR = Path("references/comtrade/hs_correlation_groups/raw")
# Optional: pass the actual inference artifact from the current Validation run.
DEFAULT_INFERENCE_PATH = None


def _hs6_cell(value) -> str | None:
    """Preserve leading zeros; reject totals, notes and non-HS6 cells."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    text = str(value).strip().replace(" ", "")
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not text.isdigit() or len(text) > 6:
        return None
    return text.zfill(6)


def _load_correlations(correlation_dir: Path) -> pd.DataFrame:
    """Read real correlation sheets, not simplified conversion tables.

    Columns are discovered from the second header row; this also handles
    the empty middle column in the HS2017-to-HS2012 workbook.
    """
    root = Path(correlation_dir)
    paths = sorted(root.glob("*.xlsx"))
    if not paths:
        raise FileNotFoundError(f"No .xlsx correlation files in {root}")
    records = []
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                if "correlation" not in ws.title.lower():
                    continue
                rows = ws.iter_rows(values_only=True)
                next(rows, None)  # 'Between' introductory row
                header = next(rows, None)
                if not header:
                    raise ValueError(f"Missing column headings: {path.name}: {ws.title}")
                normalized = [str(x or "").upper().replace(" ", "") for x in header]
                candidates = [i for i, x in enumerate(normalized)
                              if (x.startswith("HS") and "2017" in x)]
                others = [i for i, x in enumerate(normalized)
                          if x.startswith("HS") and "2017" not in x]
                rel = [i for i, x in enumerate(normalized) if "RELATIONSHIP" in x]
                if len(candidates) != 1 or len(others) != 1 or len(rel) != 1:
                    raise ValueError(f"Unrecognized correlation header: {path.name}: {ws.title}: {header}")
                hi, oi, ri = candidates[0], others[0], rel[0]
                other_version = normalized[oi]
                for row in rows:
                    h5 = _hs6_cell(row[hi] if hi < len(row) else None)
                    source = _hs6_cell(row[oi] if oi < len(row) else None)
                    if h5 is None or source is None:
                        continue
                    records.append({
                        "target_cmd": h5,
                        "related_hs_version": other_version,
                        "related_hs6": source,
                        "related_hs4": source[:4],
                        "relationship": str(row[ri] or "").strip(),
                        "workbook": path.name,
                        "worksheet": ws.title,
                    })
        finally:
            wb.close()
    if not records:
        raise ValueError(f"No HS2017 correlation relationships in {root}")
    return pd.DataFrame(records).drop_duplicates()


def _correlation_diagnostics(codes: tuple[str, ...], correlation_dir: Path):
    correlations = _load_correlations(correlation_dir)
    relevant = correlations[correlations.target_cmd.isin(codes)].copy()
    # The HS2017 code universe is obtained from reference tables, not from
    # the observed data, so missing siblings are not silently inferred.
    h5_universe = set(correlations.target_cmd)
    siblings = []
    for code in codes:
        members = sorted(c for c in h5_universe if c[:4] == code[:4] and c != code)
        siblings.append({
            "target_cmd": code,
            "target_hs4": code[:4],
            "other_h5_hs6_count": len(members),
            "other_h5_hs6": ", ".join(members),
            "h5_code_in_reference": code in h5_universe,
        })
    return relevant, pd.DataFrame(siblings)


def _inference_support(con, codes: tuple[str, ...], inference_path: Path | None):
    if inference_path is None:
        return pd.DataFrame()
    path = Path(inference_path)
    if not path.is_file():
        raise FileNotFoundError(f"Validation inference: {path}")
    cols = _columns(con, path)
    needed = {"target_cmd", "year", "hs6_seen", "hs4_seen"}
    if not needed.issubset(cols):
        raise ValueError(f"Inference missing columns: {sorted(needed - cols)}")
    prediction = ("model_prediction_quantity_kg" if "model_prediction_quantity_kg" in cols
                  else None)
    pred_expr = (f"COUNT(*) FILTER (WHERE {prediction} > 0 AND isfinite({prediction}))"
                 if prediction else "NULL")
    return _query_with_codes(con, f"""
        SELECT CAST(target_cmd AS VARCHAR) AS target_cmd, year,
               hs6_seen, hs4_seen, COUNT(*) AS rows,
               {pred_expr} AS predicted_rows
        FROM read_parquet({_sql_literal(path)})
        WHERE CAST(target_cmd AS VARCHAR) IN ({{codes}})
        GROUP BY 1, 2, 3, 4 ORDER BY 1, 2
    """, codes)


def _mapping_section(code: str, diagnostics: dict[str, pd.DataFrame]) -> list[str]:
    correlations = diagnostics.get("hs_correlation", pd.DataFrame())
    support = diagnostics.get("model_support", pd.DataFrame())
    siblings = diagnostics.get("h5_siblings", pd.DataFrame())
    lines = ["#### Official HS correlation evidence", ""]
    rel = correlations[correlations.target_cmd == code] if not correlations.empty else correlations
    if rel.empty:
        lines += ["No matching HS2017 correlation rows found.", ""]
    else:
        compact = (rel.groupby(["related_hs_version", "related_hs4", "relationship", "workbook", "worksheet"],
                               dropna=False)["related_hs6"]
                   .agg(lambda x: ", ".join(sorted(set(x)))).reset_index())
        lines += [_markdown_table(compact), ""]
    sub = siblings[siblings.target_cmd == code] if not siblings.empty else siblings
    if not sub.empty:
        lines += ["**Other HS2017 codes in this HS4:**", _markdown_table(sub), ""]
    sub = support[support.target_cmd == code] if not support.empty else support
    if sub.empty:
        lines += ["**Model HS6/HS4 support:** not checked (no inference path supplied).", ""]
    else:
        lines += ["**Model training-category support (inference flags):**", _markdown_table(sub), ""]
    lines += ["*Correlations document code relationships; training support is determined only by inference flags.*", ""]
    return lines


def _sql_literal(value: Path | str) -> str:
    return "'" + str(value).replace("'", "''") + "'"

def _normalize_codes(hs_codes: Iterable[str | int]) -> tuple[str, ...]:
    codes = tuple(dict.fromkeys(str(x).strip().zfill(6) for x in hs_codes))
    if not codes:
        raise ValueError("Provide at least one HS6 code.")

    bad = [x for x in codes if len(x) != 6 or not x.isdigit()]
    if bad:
        raise ValueError(
            "HS codes must contain exactly six digits. Invalid: "
            + ", ".join(bad)
        )

    return codes

def _columns(con: duckdb.DuckDBPyConnection, path: Path) -> set[str]:
    rows = con.execute(
        f"DESCRIBE SELECT * FROM read_parquet({_sql_literal(path)})"
    ).fetchall()

    return {str(row[0]) for row in rows}

def _require_columns(
    con: duckdb.DuckDBPyConnection,
    path: Path,
    required: set[str],
    *,
    label: str,
) -> None:

    missing = sorted(required - _columns(con, path))
    if missing:
        raise RuntimeError(
            f"{label} is missing required columns: " + ", ".join(missing)
        )

def _query_with_codes(
    con: duckdb.DuckDBPyConnection,
    sql: str,
    codes: tuple[str, ...],
) -> pd.DataFrame:

    placeholders = ", ".join(["?"] * len(codes))
    return con.execute(sql.replace("{codes}", placeholders), codes).fetchdf()

def _markdown_table(df: pd.DataFrame) -> str:
    if df.empty:
        return "_No rows._"

    view = df.copy().fillna("")

    if "year" in view.columns:
        view["year"] = view["year"].apply(
            lambda x: str(int(x)) if x != "" else ""
        )

    def fmt(value) -> str:
        if isinstance(value, bool):
            return "yes" if value else "no"

        if isinstance(value, int):
            return f"{value:,}"

        if isinstance(value, float):
            if value != value:
                return ""

            if abs(value - round(value)) < 1e-10:
                return f"{int(round(value)):,}"
            return f"{value:,.3f}"
        return str(value).replace("|", r"\\|").replace("\n", " ")

    headers = [str(c) for c in view.columns]
    rows = [[fmt(v) for v in row] for row in view.to_numpy()]

    def line(values: list[str]) -> str:
        return "| " + " | ".join(values) + " |"
    return "\n".join(
        [
            line(headers),
            line(["---"] * len(headers)),
            *[line(row) for row in rows],
        ]
    )

def _boolean_value(value) -> bool:
    """Interpret real booleans and common boolean strings correctly."""

    if value is None:
        return False
    try:
        if pd.isna(value):
            return False
    except (TypeError, ValueError):
        pass
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().lower()
    
    if text in {"true", "1", "yes", "y", "t"}:
        return True

    if text in {"false", "0", "no", "n", "f", ""}:
        return False

    raise ValueError(
        f"Could not interpret allocation_is_assumption={value!r} as boolean."
    )

def _stage_counts(
    con: duckdb.DuckDBPyConnection,
    codes: tuple[str, ...],
    stages: dict[str, Path],
) -> pd.DataFrame:

    frames: list[pd.DataFrame] = []

    for stage, path in stages.items():
        _require_columns(con, path, {"target_cmd", "year"}, label=stage)

        frame = _query_with_codes(

            con,

            f"""

            SELECT target_cmd, year, COUNT(*) AS n_rows

            FROM read_parquet({_sql_literal(path)})

            WHERE target_cmd IN ({{codes}})

            GROUP BY target_cmd, year

            ORDER BY target_cmd, year

            """,

            codes,

        )

        frame["stage"] = STAGE_LABELS[stage]

        frames.append(frame)

    return pd.concat(frames, ignore_index=True)

def _journey_wide(journey: pd.DataFrame) -> pd.DataFrame:

    if journey.empty:

        return journey.copy()

    order = [STAGE_LABELS[key] for key in STAGES]

    wide = (

        journey.pivot_table(

            index=["target_cmd", "year"],

            columns="stage",

            values="n_rows",

            aggfunc="first",

        )

        .reset_index()

    )

    for column in order:

        if column not in wide.columns:

            wide[column] = 0

        else:

            wide[column] = (

                pd.to_numeric(wide[column], errors="coerce")

                .fillna(0)

                .astype(int)

            )

    return wide[["target_cmd", "year", *order]].sort_values(

        ["target_cmd", "year"]

    )

def _hs_sync_provenance(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    path: Path,

) -> pd.DataFrame:

    required = {

        "target_cmd",

        "year",

        "source_hs_original",

        "source_cmd_original",

        "hs_sync_method",

        "allocation_strategy",

        "allocation_is_assumption",

        "allocation_weight",

        "primary_value_original",

        "primary_value_hs_synced",

    }

    _require_columns(con, path, required, label="HS-sync")

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            source_hs_original,

            source_cmd_original,

            hs_sync_method,

            allocation_strategy,

            allocation_is_assumption,

            allocation_weight,

            COUNT(*) AS n_rows,

            SUM(primary_value_original) AS original_value,

            SUM(primary_value_hs_synced) AS hs_synced_value

        FROM read_parquet({_sql_literal(path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY

            target_cmd,

            year,

            source_hs_original,

            source_cmd_original,

            hs_sync_method,

            allocation_strategy,

            allocation_is_assumption,

            allocation_weight

        ORDER BY target_cmd, year, n_rows DESC, source_cmd_original

        """,

        codes,

    )

def _quantity_summary(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    path: Path,

) -> pd.DataFrame:

    cols = _columns(con, path)

    expressions = ["COUNT(*) AS n_rows"]

    if "synchronized_quantity_kg" in cols:

        expressions.append(

            """

            COUNT(*) FILTER (

                WHERE synchronized_quantity_kg IS NOT NULL

                  AND synchronized_quantity_kg > 0

                  AND isfinite(synchronized_quantity_kg)

            ) AS n_valid_synchronized_quantity

            """

        )

    if "is_synchronized_quantity_estimated" in cols:

        expressions.append(

            """

            COUNT(*) FILTER (

                WHERE COALESCE(is_synchronized_quantity_estimated, FALSE)

            ) AS n_estimated_synchronized_quantity

            """

        )

    select_expr = ",\n".join(expressions)

    return _query_with_codes(

        con,

        f"""

        SELECT target_cmd, year, {select_expr}

        FROM read_parquet({_sql_literal(path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY target_cmd, year

        ORDER BY target_cmd, year

        """,

        codes,

    )

def _quantity_source_breakdown(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    path: Path,

) -> pd.DataFrame:

    if "synchronized_quantity_source" not in _columns(con, path):

        return pd.DataFrame()

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            COALESCE(

                CAST(synchronized_quantity_source AS VARCHAR),

                '<NULL>'

            ) AS synchronized_quantity_source,

            COUNT(*) AS n_rows

        FROM read_parquet({_sql_literal(path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY target_cmd, year, synchronized_quantity_source

        ORDER BY target_cmd, year, n_rows DESC

        """,

        codes,

    )

def _cif_fob_filter_reason_summary(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    quantity_path: Path,

) -> pd.DataFrame:

    required = {

        "target_cmd",

        "year",

        "flow_code",

        "is_reported",

        "is_aggregate",

    }

    _require_columns(

        con,

        quantity_path,

        required,

        label="Quantity input to CIF/FOB",

    )

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            COUNT(*) AS quantity_source_rows,

            COUNT(*) FILTER (

                WHERE flow_code NOT IN ('M', 'X')

            ) AS noncanonical_flow_rows,

            COUNT(*) FILTER (

                WHERE COALESCE(is_reported, FALSE) = FALSE

            ) AS not_explicitly_reported_rows,

            COUNT(*) FILTER (

                WHERE COALESCE(is_aggregate, FALSE) = TRUE

            ) AS marked_aggregate_rows,

            COUNT(*) FILTER (

                WHERE flow_code IN ('M', 'X')

                  AND COALESCE(is_reported, FALSE) = TRUE

                  AND COALESCE(is_aggregate, FALSE) = FALSE

            ) AS rows_passing_all_filters

        FROM read_parquet({_sql_literal(quantity_path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY target_cmd, year

        ORDER BY target_cmd, year

        """,

        codes,

    )

def _cif_fob_transition(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    quantity_path: Path,

    cif_fob_path: Path,

) -> pd.DataFrame:

    """

    Explain Quantity -> CIF/FOB reduction using the actual CIF/FOB

    mirror-pair construction.

    Stages:

      Quantity source rows

      -> retained reporter-side rows

      -> aggregated reporter-side rows

      -> oriented bilateral mirror pairs

      -> self-trade exclusion

      -> actual CIF/FOB rows

    """

    required = {

        "target_cmd",

        "year",

        "reporter_iso",

        "partner_iso",

        "flow_code",

        "is_reported",

        "is_aggregate",

    }

    _require_columns(

        con,

        quantity_path,

        required,

        label="Quantity input to CIF/FOB",

    )

    cif_required = {

        "target_cmd",

        "year",

        "exporter_iso",

        "importer_iso",

    }

    _require_columns(

        con,

        cif_fob_path,

        cif_required,

        label="CIF/FOB output",

    )

    placeholders = ", ".join(

        ["?"] * len(codes)

    )

    params = (

        tuple(codes)

        + tuple(codes)

    )

    return con.execute(

        f"""

        WITH classified AS (

            SELECT

                target_cmd,

                year,

                reporter_iso,

                partner_iso,

                flow_code,

                CASE

                    WHEN flow_code NOT IN ('M', 'X')

                        THEN 'excluded_noncanonical_flow'

                    WHEN COALESCE(

                        is_reported,

                        FALSE

                    ) = FALSE

                        THEN 'excluded_not_reported'

                    WHEN COALESCE(

                        is_aggregate,

                        FALSE

                    ) = TRUE

                        THEN 'excluded_aggregate'

                    ELSE 'retained'

                END

                    AS diagnostic_fate

            FROM read_parquet(

                {_sql_literal(quantity_path)}

            )

            WHERE target_cmd IN (

                {placeholders}

            )

        ),

        source_summary AS (

            SELECT

                target_cmd,

                year,

                COUNT(*)

                    AS quantity_source_rows,

                COUNT(*) FILTER (

                    WHERE diagnostic_fate = 'retained'

                )

                    AS retained_source_rows,

                COUNT(

                    DISTINCT (

                        reporter_iso,

                        partner_iso,

                        flow_code

                    )

                ) FILTER (

                    WHERE diagnostic_fate = 'retained'

                )

                    AS canonical_reporter_rows

            FROM classified

            GROUP BY

                target_cmd,

                year

        ),

        retained AS (

            SELECT DISTINCT

                target_cmd,

                year,

                reporter_iso,

                partner_iso,

                flow_code

            FROM classified

            WHERE diagnostic_fate = 'retained'

        ),

        oriented AS (

            SELECT

                target_cmd,

                year,

                reporter_iso

                    AS exporter_iso,

                partner_iso

                    AS importer_iso,

                'X'

                    AS source_side

            FROM retained

            WHERE flow_code = 'X'

            UNION ALL

            SELECT

                target_cmd,

                year,

                partner_iso

                    AS exporter_iso,

                reporter_iso

                    AS importer_iso,

                'M'

                    AS source_side

            FROM retained

            WHERE flow_code = 'M'

        ),

        pair_summary AS (

            SELECT

                target_cmd,

                year,

                COUNT(

                    DISTINCT (

                        exporter_iso,

                        importer_iso

                    )

                )

                    AS oriented_mirror_pairs,

                COUNT(

                    DISTINCT (

                        exporter_iso,

                        importer_iso

                    )

                ) FILTER (

                    WHERE exporter_iso = importer_iso

                )

                    AS self_trade_pairs,

                COUNT(

                    DISTINCT (

                        exporter_iso,

                        importer_iso

                    )

                ) FILTER (

                    WHERE exporter_iso <> importer_iso

                )

                    AS expected_mirror_pairs

            FROM oriented

            GROUP BY

                target_cmd,

                year

        ),

        actual AS (

            SELECT

                target_cmd,

                year,

                COUNT(*)

                    AS actual_cif_fob_rows

            FROM read_parquet(

                {_sql_literal(cif_fob_path)}

            )

            WHERE target_cmd IN (

                {placeholders}

            )

            GROUP BY

                target_cmd,

                year

        )

        SELECT

            s.target_cmd,

            s.year,

            s.quantity_source_rows,

            s.retained_source_rows,

            s.canonical_reporter_rows,

            p.oriented_mirror_pairs,

            p.self_trade_pairs,

            p.expected_mirror_pairs,

            COALESCE(

                a.actual_cif_fob_rows,

                0

            )

                AS actual_cif_fob_rows,

            COALESCE(

                a.actual_cif_fob_rows,

                0

            )

            -

            p.expected_mirror_pairs

                AS unexplained_difference

        FROM source_summary AS s

        LEFT JOIN pair_summary AS p

            USING (

                target_cmd,

                year

            )

        LEFT JOIN actual AS a

            USING (

                target_cmd,

                year

            )

        ORDER BY

            s.target_cmd,

            s.year

        """,

        params,

    ).fetchdf()

def _cif_fob_exclusion_patterns(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    quantity_path: Path,

) -> pd.DataFrame:

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            flow_code,

            COALESCE(partner_iso, '<NULL>') AS partner_iso,

            flow_code NOT IN ('M', 'X') AS noncanonical_flow,

            COALESCE(is_reported, FALSE) = FALSE AS not_explicitly_reported,

            COALESCE(is_aggregate, FALSE) = TRUE AS marked_aggregate,

            COUNT(*) AS n_source_rows

        FROM read_parquet({_sql_literal(quantity_path)})

        WHERE target_cmd IN ({{codes}})

          AND (

                flow_code NOT IN ('M', 'X')

             OR COALESCE(is_reported, FALSE) = FALSE

             OR COALESCE(is_aggregate, FALSE) = TRUE

          )

        GROUP BY

            target_cmd,

            year,

            flow_code,

            partner_iso,

            noncanonical_flow,

            not_explicitly_reported,

            marked_aggregate

        ORDER BY target_cmd, year, n_source_rows DESC, flow_code, partner_iso

        """,

        codes,

    )

def _cif_fob_aggregation_multiplicity(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    quantity_path: Path,

) -> pd.DataFrame:

    return _query_with_codes(

        con,

        f"""

        WITH retained AS (

            SELECT

                target_cmd,

                year,

                reporter_iso,

                partner_iso,

                flow_code,

                COUNT(*) AS source_components

            FROM read_parquet({_sql_literal(quantity_path)})

            WHERE target_cmd IN ({{codes}})

              AND flow_code IN ('M', 'X')

              AND COALESCE(is_reported, FALSE) = TRUE

              AND COALESCE(is_aggregate, FALSE) = FALSE

            GROUP BY

                target_cmd,

                year,

                reporter_iso,

                partner_iso,

                flow_code

        )

        SELECT

            target_cmd,

            year,

            source_components,

            COUNT(*) AS n_canonical_groups

        FROM retained

        GROUP BY target_cmd, year, source_components

        ORDER BY target_cmd, year, source_components

        """,

        codes,

    )

def _simple_transition(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    left_path: Path,

    right_path: Path,

    *,

    left_label: str,

    right_label: str,

) -> pd.DataFrame:

    placeholders = ", ".join(["?"] * len(codes))

    return con.execute(

        f"""

        WITH left_stage AS (

            SELECT target_cmd, year, COUNT(*) AS left_rows

            FROM read_parquet({_sql_literal(left_path)})

            WHERE target_cmd IN ({placeholders})

            GROUP BY target_cmd, year

        ),

        right_stage AS (

            SELECT target_cmd, year, COUNT(*) AS right_rows

            FROM read_parquet({_sql_literal(right_path)})

            WHERE target_cmd IN ({placeholders})

            GROUP BY target_cmd, year

        )

        SELECT

            COALESCE(l.target_cmd, r.target_cmd) AS target_cmd,

            COALESCE(l.year, r.year) AS year,

            COALESCE(l.left_rows, 0) AS "{left_label}",

            COALESCE(r.right_rows, 0) AS "{right_label}",

            COALESCE(r.right_rows, 0) - COALESCE(l.left_rows, 0)

                AS row_change

        FROM left_stage AS l

        FULL OUTER JOIN right_stage AS r USING (target_cmd, year)

        ORDER BY target_cmd, year

        """,

        tuple(codes) + tuple(codes),

    ).fetchdf()

def _validation_support(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    path: Path,

) -> pd.DataFrame:

    required = {

        "target_cmd",

        "year",

        "quantity_selected_kg",

        "trade_value_reconciled",

        "export_synchronized_quantity_kg",

        "import_synchronized_quantity_kg",

        "export_is_synchronized_quantity_estimated",

        "import_is_synchronized_quantity_estimated",

    }

    _require_columns(con, path, required, label="Validation quantity selection")

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            COUNT(*) AS n_rows,

            COUNT(*) FILTER (

                WHERE quantity_selected_kg IS NOT NULL

                  AND quantity_selected_kg > 0

                  AND isfinite(quantity_selected_kg)

            ) AS n_selected_quantity,

            COUNT(*) FILTER (

                WHERE export_synchronized_quantity_kg IS NOT NULL

                  AND import_synchronized_quantity_kg IS NOT NULL

                  AND export_synchronized_quantity_kg > 0

                  AND import_synchronized_quantity_kg > 0

                  AND isfinite(export_synchronized_quantity_kg)

                  AND isfinite(import_synchronized_quantity_kg)

            ) AS n_bilateral_quantity,

            COUNT(*) FILTER (

                WHERE export_synchronized_quantity_kg IS NOT NULL

                  AND import_synchronized_quantity_kg IS NOT NULL

                  AND export_synchronized_quantity_kg > 0

                  AND import_synchronized_quantity_kg > 0

                  AND isfinite(export_synchronized_quantity_kg)

                  AND isfinite(import_synchronized_quantity_kg)

                  AND COALESCE(

                      export_is_synchronized_quantity_estimated,

                      FALSE

                  ) = FALSE

                  AND COALESCE(

                      import_is_synchronized_quantity_estimated,

                      FALSE

                  ) = FALSE

            ) AS n_direct_bilateral,

            COUNT(*) FILTER (

                WHERE export_synchronized_quantity_kg IS NOT NULL

                  AND import_synchronized_quantity_kg IS NOT NULL

                  AND export_synchronized_quantity_kg > 0

                  AND import_synchronized_quantity_kg > 0

                  AND isfinite(export_synchronized_quantity_kg)

                  AND isfinite(import_synchronized_quantity_kg)

                  AND COALESCE(

                      export_is_synchronized_quantity_estimated,

                      FALSE

                  ) = FALSE

                  AND COALESCE(

                      import_is_synchronized_quantity_estimated,

                      FALSE

                  ) = FALSE

                  AND quantity_selected_kg IS NOT NULL

                  AND quantity_selected_kg > 0

                  AND isfinite(quantity_selected_kg)

                  AND trade_value_reconciled IS NOT NULL

                  AND trade_value_reconciled > 0

                  AND isfinite(trade_value_reconciled)

            ) AS n_reference_eligible

        FROM read_parquet({_sql_literal(path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY target_cmd, year

        ORDER BY target_cmd, year

        """,

        codes,

    )

def _validation_quantity_sources(

    con: duckdb.DuckDBPyConnection,

    codes: tuple[str, ...],

    path: Path,

) -> pd.DataFrame:

    needed = {

        "target_cmd",

        "year",

        "quantity_selected_candidate",

        "quantity_selected_source",

    }

    if not needed.issubset(_columns(con, path)):

        return pd.DataFrame()

    return _query_with_codes(

        con,

        f"""

        SELECT

            target_cmd,

            year,

            COALESCE(

                CAST(quantity_selected_candidate AS VARCHAR),

                '<NULL>'

            ) AS quantity_selected_candidate,

            COALESCE(

                CAST(quantity_selected_source AS VARCHAR),

                '<NULL>'

            ) AS quantity_selected_source,

            COUNT(*) AS n_rows

        FROM read_parquet({_sql_literal(path)})

        WHERE target_cmd IN ({{codes}})

        GROUP BY

            target_cmd,

            year,

            quantity_selected_candidate,

            quantity_selected_source

        ORDER BY target_cmd, year, n_rows DESC

        """,

        codes,

    )

def _build_conclusion(

    *,

    assumption_based: bool,

    deterministic_mapping: bool,

    hs_sync_rows: int,

    quantity_rows: int,

    cif_unexplained_difference: int,

    bilateral: int,

    direct_bilateral: int,

    reference_eligible: int,

) -> str:

    """

    Build a short evidence-based conclusion for one HS6 code.

    Every statement is derived from the diagnostics collected for that code.

    No specific HS code or expected failure mode is hard-coded.

    """

    parts: list[str] = []

    # ----------------------------------------------------------

    # HS-sync

    # ----------------------------------------------------------

    if assumption_based:

        parts.append(

            "HS-sync uses a mapping that is marked as assumption-based"

        )

    elif deterministic_mapping:

        parts.append(

            "HS-sync uses a deterministic mapping"

        )

    else:

        parts.append(

            "HS-sync does not classify the mapping as either "

            "deterministic or assumption-based"

        )

    # ----------------------------------------------------------

    # Quantity

    # ----------------------------------------------------------

    if quantity_rows == hs_sync_rows:

        parts.append(

            "Quantity preserves the HS-sync row count"

        )

    elif quantity_rows < hs_sync_rows:

        parts.append(

            f"The row count decreases from {hs_sync_rows:,} "

            f"to {quantity_rows:,} in Quantity"

        )

    else:

        parts.append(

            f"The row count increases from {hs_sync_rows:,} "

            f"to {quantity_rows:,} in Quantity"

        )

    # ----------------------------------------------------------

    # CIF/FOB

    # ----------------------------------------------------------

    if cif_unexplained_difference == 0:

        parts.append(

            "The CIF/FOB row-count change is fully explained "

            "by its filtering and reporter-side aggregation rules"

        )

    else:

        parts.append(

            f"CIF/FOB has an unexplained row-count difference of "

            f"{cif_unexplained_difference:,}"

        )

    # ----------------------------------------------------------

    # Validation

    # ----------------------------------------------------------

    if reference_eligible > 0:

        parts.append(

            f"Validation contains {reference_eligible:,} "

            "reference-eligible observation(s)"

        )

    elif bilateral == 0:

        parts.append(

            "Validation contains no observations with usable "

            "quantities on both trade sides"

        )

    elif direct_bilateral == 0:

        parts.append(

            "Validation contains bilateral quantity observations, "

            "but none are direct non-estimated bilateral observations"

        )

    else:

        parts.append(

            "Validation contains direct bilateral quantity evidence, "

            "but no observations satisfy the full reference-eligibility rule"

        )

    return ". ".join(parts) + "."

def collect_hs6_diagnostics(

    hs_codes: Iterable[str | int],

    *,

    stages: dict[str, Path] = STAGES,

    con: duckdb.DuckDBPyConnection | None = None,
    correlation_dir: Path = DEFAULT_CORRELATION_DIR,
    inference_path: Path | None = DEFAULT_INFERENCE_PATH,

) -> dict[str, pd.DataFrame]:

    codes = _normalize_codes(hs_codes)

    owns_connection = con is None

    if con is None:

        con = duckdb.connect()

    try:

        for stage, path in stages.items():

            if not path.is_file():

                raise FileNotFoundError(f"{stage}: {path}")

        hs_correlation, h5_siblings = _correlation_diagnostics(codes, correlation_dir)
        result = {

            "journey": _stage_counts(con, codes, stages),

            "hs_sync_provenance": _hs_sync_provenance(

                con, codes, stages["hs_sync"]

            ),

            "quantity_summary": _quantity_summary(

                con, codes, stages["quantity"]

            ),

            "quantity_source_breakdown": _quantity_source_breakdown(

                con, codes, stages["quantity"]

            ),

            "cif_fob_filter_reason_summary": _cif_fob_filter_reason_summary(

                con, codes, stages["quantity"]

            ),

            "cif_fob_transition": _cif_fob_transition(

                con, codes, stages["quantity"], stages["cif_fob"]

            ),

            "cif_fob_exclusion_patterns": _cif_fob_exclusion_patterns(

                con, codes, stages["quantity"]

            ),

            "cif_fob_aggregation_multiplicity": (

                _cif_fob_aggregation_multiplicity(

                    con, codes, stages["quantity"]

                )

            ),

            "cif_fob_to_trade_sync": _simple_transition(

                con,

                codes,

                stages["cif_fob"],

                stages["trade_sync"],

                left_label="cif_fob_rows",

                right_label="trade_sync_rows",

            ),

            "trade_sync_to_validation": _simple_transition(

                con,

                codes,

                stages["trade_sync"],

                stages["validation_prepared"],

                left_label="trade_sync_rows",

                right_label="validation_prepared_rows",

            ),

            "validation_prepared_to_quantity": _simple_transition(

                con,

                codes,

                stages["validation_prepared"],

                stages["validation_quantity"],

                left_label="validation_prepared_rows",

                right_label="validation_quantity_rows",

            ),

            "validation_support": _validation_support(

                con, codes, stages["validation_quantity"]

            ),

            "validation_quantity_sources": _validation_quantity_sources(

                con, codes, stages["validation_quantity"]

            ),

        }
        result["hs_correlation"] = hs_correlation
        result["h5_siblings"] = h5_siblings
        result["model_support"] = _inference_support(con, codes, inference_path)
        return result

    finally:

        if owns_connection:

            con.close()

def _transition_changed(frame: pd.DataFrame, code: str) -> bool:

    local = frame[frame["target_cmd"] == code]

    return (not local.empty) and bool((local["row_change"] != 0).any())

def _code_story(

    code: str,

    diagnostics: dict[str, pd.DataFrame],

) -> list[str]:

    lines = [f"### HS6 `{code}`", ""]

    # Values collected from the diagnostics and used for the

    # evidence-based conclusion at the end.

    assumption_based = False

    deterministic_mapping = False

    hs_sync_rows = 0

    quantity_rows = 0

    cif_unexplained_difference = 0

    bilateral = 0

    direct_bilateral = 0

    reference_eligible = 0

    journey = diagnostics["journey"]

    local_journey = journey[journey["target_cmd"] == code]

    if local_journey.empty:

        return lines + ["This code was not found in the configured stages.", ""]

    years = sorted(local_journey["year"].dropna().astype(int).unique())

    if years:

        lines += [

            f"This code appears in the pipeline for **{years[0]}–{years[-1]}**.",

            "",

        ]

    provenance = diagnostics["hs_sync_provenance"]

    provenance = provenance[provenance["target_cmd"] == code]

    if not provenance.empty:

        source_codes = ", ".join(

            sorted(

                set(

                    provenance["source_cmd_original"]

                    .dropna()

                    .astype(str)

                )

            )

        )

        method_set = set(

            provenance["hs_sync_method"]

            .dropna()

            .astype(str)

        )

        methods = ", ".join(

            sorted(method_set)

        )

        weights = sorted(

            set(

                float(x)

                for x in provenance["allocation_weight"].dropna()

            )

        )

        weight_text = ", ".join(

            f"{x:g}" for x in weights

        ) or "n/a"

        assumption_based = any(

            _boolean_value(x)

            for x in provenance[

                "allocation_is_assumption"

            ].tolist()

        )

        deterministic_mapping = (

            not assumption_based

            and bool(method_set)

            and all(

                method.startswith("deterministic")

                for method in method_set

            )

        )

        assumption_text = (

            "assumption-based"

            if assumption_based

            else "not assumption-based"

        )

        lines += [

            (

                "**1. HS-sync origin.** "

                f"Source code(s): **{source_codes}**. "

                f"Mapping method(s): **{methods}**. "

                f"Allocation weight(s): **{weight_text}**. "

                f"The mapping is **{assumption_text}**."

            ),

            "",

        ]

    hs_sync_rows = int(

        local_journey.loc[

            local_journey["stage"] == "HS-sync",

            "n_rows",

        ].sum()

    )

    quantity_rows = int(

        local_journey.loc[

            local_journey["stage"] == "Quantity",

            "n_rows",

        ].sum()

    )

    if hs_sync_rows == quantity_rows:

        quantity_text = (

            f"Quantity keeps all **{quantity_rows:,}** HS-sync row(s). "

            "There is no row-count reduction here."

        )

    else:

        quantity_text = (

            f"The row count changes from **{hs_sync_rows:,}** to "

            f"**{quantity_rows:,}** here. This needs inspection."

        )

    lines += [f"**2. Quantity.** {quantity_text}", ""]

    transition = diagnostics["cif_fob_transition"]

    transition = transition[transition["target_cmd"] == code]

    reasons = diagnostics["cif_fob_filter_reason_summary"]

    reasons = reasons[reasons["target_cmd"] == code]

    patterns = diagnostics["cif_fob_exclusion_patterns"]

    patterns = patterns[patterns["target_cmd"] == code]

    if not transition.empty:

        source_rows = int(transition["quantity_source_rows"].sum())

        retained_rows = int(transition["retained_source_rows"].sum())

        canonical_reporter_rows = int(

            transition[

                "canonical_reporter_rows"

            ].sum()

        )

        oriented_mirror_pairs = int(

            transition[

                "oriented_mirror_pairs"

            ].sum()

        )

        self_trade_pairs = int(

            transition[

                "self_trade_pairs"

            ].sum()

        )

        expected_rows = int(

            transition[

                "expected_mirror_pairs"

            ].sum()

        )

        actual_rows = int(

            transition[

                "actual_cif_fob_rows"

            ].sum()

        )        

        actual_rows = int(transition["actual_cif_fob_rows"].sum())

        cif_unexplained_difference = int(

                transition[

                    "unexplained_difference"

                ].abs().sum()

            )

        noncanonical = int(reasons["noncanonical_flow_rows"].sum())

        not_reported = int(reasons["not_explicitly_reported_rows"].sum())

        aggregate = int(reasons["marked_aggregate_rows"].sum())

        flow_codes = sorted(

            set(

                patterns.loc[

                    patterns["noncanonical_flow"] == True,

                    "flow_code",

                ].astype(str)

            )

        )

        flow_note = ""

        if flow_codes:

            flow_note = " (" + ", ".join(f"`{x}`" for x in flow_codes) + ")"

        lines += [

            (

                "**3. Quantity → CIF/FOB.** "

                f"CIF/FOB receives **{source_rows:,}** source/component row(s). "

                f"Among them, **{noncanonical:,}** have a noncanonical flow"

                f"{flow_note}, **{not_reported:,}** are not explicitly "

                f"reported, and **{aggregate:,}** are marked aggregate upstream."

            ),

            "",

            (

                "These reason counts **can overlap**. A row can be both "

                "not explicitly reported and marked aggregate, so do not add "

                "the three reason counts together."

            ),

            "",

            (

                f"After all CIF/FOB input rules are applied, "

                f"**{retained_rows:,}** source/component row(s) remain. "

                f"They aggregate to **{canonical_reporter_rows:,}** canonical "

                "reporter-side row(s) using "

                "`year × reporter × partner × flow × target_cmd`."

            ),

            "",

            (

                f"Those reporter-side rows are then oriented to bilateral "

                f"`year × exporter × importer × target_cmd` keys. "

                f"This produces **{oriented_mirror_pairs:,}** candidate mirror "

                f"pair(s), of which **{self_trade_pairs:,}** are self-trade "

                f"pairs and are excluded."

            ),

            "",

            (

                f"The expected CIF/FOB population is therefore "

                f"**{expected_rows:,}** bilateral mirror-pair row(s)."

            ),

            "",

        ]

        if cif_unexplained_difference == 0:

            lines += [

                (

                    f"The actual CIF/FOB output also has **{actual_rows:,}** "

                    "row(s). The row-count change is therefore **fully "

                    "accounted for** by the configured filters and aggregation. "

                    "There is no unexplained row loss at this boundary."

                ),

                "",

            ]

        else:

            lines += [

                (

                    f"The actual CIF/FOB output has **{actual_rows:,}** row(s), "

                    f"leaving an **unexplained difference of {cif_unexplained_difference:,}**. "

                    "This should be investigated."

                ),

                "",

            ]

    changed = []

    for key, left_name, right_name in [

        ("cif_fob_to_trade_sync", "CIF/FOB", "Trade Sync"),

        ("trade_sync_to_validation", "Trade Sync", "Validation prepared"),

        (

            "validation_prepared_to_quantity",

            "Validation prepared",

            "Validation quantity selection",

        ),

    ]:

        if _transition_changed(diagnostics[key], code):

            changed.append(f"{left_name} → {right_name}")

    if changed:

        downstream_text = (

            "Further row-count changes occur at " + ", ".join(changed) + "."

        )

    else:

        downstream_text = (

            "After CIF/FOB, the row count stays unchanged through Trade Sync, "

            "Validation preparation, and Validation quantity selection."

        )

    lines += [f"**4. Downstream stages.** {downstream_text}", ""]

    support = diagnostics["validation_support"]

    support = support[support["target_cmd"] == code]

    selected_sources = diagnostics["validation_quantity_sources"]

    selected_sources = selected_sources[selected_sources["target_cmd"] == code]

    if not support.empty:

        total = int(support["n_rows"].sum())

        selected = int(support["n_selected_quantity"].sum())

        bilateral = int(

            support["n_bilateral_quantity"].sum()

        )

        direct_bilateral = int(

            support["n_direct_bilateral"].sum()

        )

        reference_eligible = int(

            support["n_reference_eligible"].sum()

        )

        source_parts: list[str] = []

        if not selected_sources.empty:

            for source, group in selected_sources.groupby(

                "quantity_selected_source"

            ):

                source_parts.append(

                    f"{int(group['n_rows'].sum()):,} `{source}`"

                )

        lines += [

            (

                "**5. Validation outcome.** "

                f"Validation contains **{total:,}** row(s), and "

                f"**{selected:,}** have a selected quantity."

            ),

            "",

        ]

        if source_parts:

            lines += [

                "Selected-quantity sources: " + ", ".join(source_parts) + ".",

                "",

            ]

        lines += [

            (

                f"Of the {total:,} Validation row(s), **{bilateral:,}** have "

                f"usable quantities on both trade sides, "

                f"**{direct_bilateral:,}** have "

                "direct non-estimated bilateral quantities, and "

                f"**{reference_eligible:,}** satisfy the full "

                "reference-eligibility rule."

            ),

            "",

        ]

        if reference_eligible == 0:

            result = (

                "**Result:** this HS6 has **no reference-eligible observations** "

                "in the current Validation population."

            )

        else:

            result = (

                f"**Result:** this HS6 has **{reference_eligible:,} reference-eligible "

                "observation(s)** in the current Validation population."

            )

        lines += [result, ""]

    # ----------------------------------------------------------

    # Evidence-based overall conclusion

    # ----------------------------------------------------------

    conclusion = _build_conclusion(

        assumption_based=assumption_based,

        deterministic_mapping=deterministic_mapping,

        hs_sync_rows=hs_sync_rows,

        quantity_rows=quantity_rows,

        cif_unexplained_difference=(

            cif_unexplained_difference

        ),

        bilateral=bilateral,

        direct_bilateral=direct_bilateral,

        reference_eligible=reference_eligible,

    )

    lines += _mapping_section(code, diagnostics)
    lines += [

        "### In short",

        "",

        conclusion,

        "",

        "---",

        "",

    ]

    return lines

def _transition_note(

    frame: pd.DataFrame,

    left_name: str,

    right_name: str,

) -> str:

    if frame.empty:

        return f"No rows are available for {left_name} → {right_name}."

    changed = frame[frame["row_change"] != 0]

    if changed.empty:

        return (

            f"No row-count change is observed from {left_name} to {right_name} "

            "for the requested HS6 code(s)."

        )

    return (

        f"Row counts change from {left_name} to {right_name} for "

        f"{len(changed)} HS6-year combination(s)."

    )

def render_hs6_diagnostic_markdown(

    hs_codes: Iterable[str | int],

    diagnostics: dict[str, pd.DataFrame],

    *,

    stages: dict[str, Path] = STAGES,

) -> str:

    codes = _normalize_codes(hs_codes)

    now = datetime.now().astimezone()

    lines: list[str] = [

        "# VACODA HS6 lineage diagnostic",

        "",

        f"Generated: `{now.isoformat(timespec='seconds')}`",

        "",

        "HS6 code(s): " + ", ".join(f"`{x}`" for x in codes),

        "",

        "## What this report answers",

        "",

        "This report follows each requested HS6 code from HS-sync to Validation.",

        "",

        "1. Where did the HS6 code come from?",

        "2. At which stage did its row count change, and why?",

        "3. Is every row-count change explained by the pipeline rules?",

        "4. Why does the code have, or not have, Validation reference support?",

        "",

        "### Three terms used in the report",

        "",

        "- **Source/component row:** an upstream row before CIF/FOB aggregation.",

        (

            "- **Reporter-side trade row:** one CIF/FOB row for "

            "`year × reporter × partner × flow × target_cmd`."

        ),

        (

            "- **Reference-eligible row:** a Validation row with valid direct, "

            "non-estimated quantities on both trade sides, plus a valid selected "

            "quantity and reconciled trade value."

        ),

        "",

        (

            "**Important:** a lower row count does not automatically mean data "

            "was lost. CIF/FOB intentionally filters rows outside its canonical "

            "population and then aggregates source components."

        ),

        "",

        "## 1. Diagnosis by HS6 code",

        "",

    ]

    for code in codes:

        lines.extend(_code_story(code, diagnostics))

    lines += [

        "## 2. Pipeline row-count overview",

        "",

        (

            "One row is shown per HS6/year. Section 5 explains the CIF/FOB "

            "row-count change in detail."

        ),

        "",

        _markdown_table(_journey_wide(diagnostics["journey"])),

        "",

        "## 3. HS-sync origin",

        "",

        (

            "This table shows the source HS code, mapping method, allocation "

            "weight, and whether the mapping was marked as assumption-based."

        ),

        "",

        _markdown_table(diagnostics["hs_sync_provenance"]),

        "",

        "## 4. Quantity",

        "",

        (

            "Quantity consumes the HS-sync result and chooses/reconciles the "

            "physical quantity used downstream."

        ),

        "",

        _markdown_table(diagnostics["quantity_summary"]),

    ]

    if not diagnostics["quantity_source_breakdown"].empty:

        lines += [

            "",

            "### 4.1. Synchronized quantity source",

            "",

            _markdown_table(diagnostics["quantity_source_breakdown"]),

        ]

    lines += [

        "",

        "## 5. Quantity → CIF/FOB",

        "",

        "### 5.1. Which CIF/FOB input rules are triggered?",

        "",

        (

            "The three reason counts below are **independent and can overlap**. "

            "Do not add them together."

        ),

        "",

        _markdown_table(diagnostics["cif_fob_filter_reason_summary"]),

        "",

        "### 5.2. Row-count accounting",

        "",

        (

            "`retained_source_rows` counts Quantity rows that pass the "

            "CIF/FOB input rules. `canonical_reporter_rows` is the number "

            "of distinct reporter-side trade rows after canonical "

            "aggregation. These rows are then oriented into exporter/importer "

            "mirror pairs. `self_trade_pairs` are excluded because exporter "

            "and importer are identical. `expected_mirror_pairs` is therefore "

            "the expected final CIF/FOB row count. "

            "`unexplained_difference = 0` means the actual CIF/FOB output "

            "is exactly accounted for."

        ),

        "",

        _markdown_table(diagnostics["cif_fob_transition"]),

        "",

        "### 5.3. Excluded row patterns",

        "",

        (

            "The boolean columns show every applicable reason for a pattern. "

            "This makes overlaps visible instead of assigning only one reason."

        ),

        "",

        _markdown_table(diagnostics["cif_fob_exclusion_patterns"]),

        "",

        "### 5.4. Reporter-side reporter-side aggregation",

        "",

        (

            "`source_components` shows how many retained Quantity rows were "

            "combined into one aggregated reporter-side trade row before "

            "export/import orientation and mirror pairing."

        ),

        "",

        _markdown_table(diagnostics["cif_fob_aggregation_multiplicity"]),

        "",

        "## 6. Downstream row-count checks",

        "",

        "### 6.1. CIF/FOB → Trade Sync",

        "",

        _transition_note(

            diagnostics["cif_fob_to_trade_sync"], "CIF/FOB", "Trade Sync"

        ),

        "",

        _markdown_table(diagnostics["cif_fob_to_trade_sync"]),

        "",

        "### 6.2. Trade Sync → Validation prepared",

        "",

        _transition_note(

            diagnostics["trade_sync_to_validation"],

            "Trade Sync",

            "Validation prepared",

        ),

        "",

        _markdown_table(diagnostics["trade_sync_to_validation"]),

        "",

        "### 6.3. Validation prepared → quantity selection",

        "",

        _transition_note(

            diagnostics["validation_prepared_to_quantity"],

            "Validation prepared",

            "Validation quantity selection",

        ),

        "",

        _markdown_table(diagnostics["validation_prepared_to_quantity"]),

        "",

        "## 7. Final Validation support",

        "",

        (

            "`n_bilateral_quantity` requires usable export and import quantities. "

            "`n_direct_bilateral` additionally excludes estimated quantities. "

            "`n_reference_eligible` applies the full current reference rule."

        ),

        "",

        _markdown_table(diagnostics["validation_support"]),

    ]

    if not diagnostics["validation_quantity_sources"].empty:

        lines += [

            "",

            "### 7.1. Selected quantity source",

            "",

            (

                "`import_only` or `export_only` means Validation could select a "

                "quantity, but only one trade side supplied usable quantity "

                "evidence. Such a row is not bilateral reference evidence."

            ),

            "",

            _markdown_table(diagnostics["validation_quantity_sources"]),

        ]

    artifact_rows = pd.DataFrame(

        [

            {

                "stage": STAGE_LABELS[stage],

                "artifact": str(path),

            }

            for stage, path in stages.items()

        ]

    )

    lines += [

        "",

        "## 8. Data files used",

        "",

        (

            "These are the files read by this diagnostic so another VACODA user "

            "can reproduce the trace."

        ),

        "",

        _markdown_table(artifact_rows),

        "",

    ]

    return "\n".join(lines)

def write_hs6_diagnostic_report(

    hs_codes: Iterable[str | int],

    *,

    output_path: Path | str | None = None,

    stages: dict[str, Path] = STAGES,
    correlation_dir: Path = DEFAULT_CORRELATION_DIR,
    inference_path: Path | None = DEFAULT_INFERENCE_PATH,

) -> Path:

    codes = _normalize_codes(hs_codes)

    diagnostics = collect_hs6_diagnostics(
        codes, stages=stages, correlation_dir=correlation_dir, inference_path=inference_path
    )

    markdown = render_hs6_diagnostic_markdown(

        codes,

        diagnostics,

        stages=stages,

    )

    if output_path is None:

        DEFAULT_REPORT_DIR.mkdir(parents=True, exist_ok=True)

        output_path = DEFAULT_REPORT_DIR / (

            "hs6_" + "_".join(codes) + "_lineage_diagnostic.md"

        )

    else:

        output_path = Path(output_path)

        output_path.parent.mkdir(parents=True, exist_ok=True)

    output_path.write_text(markdown, encoding="utf-8")

    return output_path

def _parse_args() -> argparse.Namespace:

    parser = argparse.ArgumentParser(

        description=(

            "Trace one or more HS6 codes from HS-sync through Validation and "

            "write an easy-to-read Markdown diagnostic report."

        )

    )

    parser.add_argument(

        "hs_codes",

        nargs="+",

        help="One or more six-digit HS target codes.",

    )

    parser.add_argument(

        "--output",

        type=Path,

        default=None,

        help=(

            "Optional Markdown output path. By default the report is written "

            "under /storage/vacoda/output/reports/diagnostics/hs_journey/."

        ),

    )

    parser.add_argument("--correlation-dir", type=Path, default=DEFAULT_CORRELATION_DIR,
                        help="Directory with the official correlation .xlsx workbooks")
    parser.add_argument("--inference", type=Path, default=None,
                        help="Optional Validation inference parquet with hs6_seen/hs4_seen flags")
    return parser.parse_args()

def main() -> None:

    args = _parse_args()

    report_path = write_hs6_diagnostic_report(

        args.hs_codes,

        output_path=args.output,
        correlation_dir=args.correlation_dir,
        inference_path=args.inference,

    )

    print(f"HS6 diagnostic report written: {report_path}")

if __name__ == "__main__":

    main()
